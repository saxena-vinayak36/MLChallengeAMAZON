from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

OUT = "Amazon_ML_Challenge_Report.docx"

BLUE = "2E74B5"
DARK = "1F4D78"
LIGHT = "E8EEF5"
GRAY = "F2F4F7"
MUTED = "5B6573"
BLACK = "111111"

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def set_cell_width(cell, width):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'), str(width)); tcW.set(qn('w:type'), 'dxa')

def set_table_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in('w:tblW')
    if tblW is None:
        tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths))); tblW.set(qn('w:type'), 'dxa')
    tblInd = OxmlElement('w:tblInd'); tblInd.set(qn('w:w'), '120'); tblInd.set(qn('w:type'), 'dxa'); tblPr.append(tblInd)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths): col.set(qn('w:w'), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width); set_cell_margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_font(run, size=11, bold=None, color=BLACK, italic=None):
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic

def add_text(doc, text, style=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead); set_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):]); set_font(r)
    else:
        r = p.add_run(text); set_font(r)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text); set_font(r)
    return p

def add_number(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text); set_font(r)
    return p

def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = 'Table Grid'
    set_table_widths(table, widths)
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text); set_font(r, size=9.5, bold=True, color=DARK)
    for values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(text)); set_font(r, size=9.3)
    set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def h1(doc, text):
    p = doc.add_paragraph(style='Heading 1'); p.paragraph_format.keep_with_next = True
    p.clear(); r = p.add_run(text); set_font(r, 16, True, BLUE); return p

def h2(doc, text):
    p = doc.add_paragraph(style='Heading 2'); p.paragraph_format.keep_with_next = True
    p.clear(); r = p.add_run(text); set_font(r, 13, True, BLUE); return p

def page_break(doc): doc.add_page_break()

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(0.85)
sec.left_margin = sec.right_margin = Inches(1.0)
sec.header_distance = sec.footer_distance = Inches(0.49)

# Styles
normal = doc.styles['Normal']; normal.font.name = 'Calibri'; normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri'); normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.10
for name, size, before, after, color in [('Heading 1',16,16,8,BLUE), ('Heading 2',13,12,6,BLUE)]:
    st = doc.styles[name]; st.font.name='Calibri'; st._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri'); st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color); st.font.bold=True; st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)

# Header/footer
header = sec.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run('AMAZON ML CHALLENGE 2026  |  WINNING SUBMISSION REPORT'); set_font(r, 8.5, True, MUTED)
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('Business Entity Resolution Challenge'); set_font(r, 8.5, False, MUTED)

# Cover
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(34); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('AMAZON ML CHALLENGE 2026'); set_font(r, 12, True, BLUE)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(8)
r=p.add_run('Business Entity Resolution\nSolution Report'); set_font(r, 28, True, DARK)
p = doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(26)
r=p.add_run('What we built, how it ran on AWS, and what to demonstrate'); set_font(r, 13, False, MUTED)
add_table(doc, ['Challenge', 'Team', 'Report date'], [['Business Entity Resolution', 'WinKing', '30 September 2026']], [3600, 2880, 2880])
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(20)
r=p.add_run('One-line outcome: '); set_font(r, 11, True, DARK)
r=p.add_run('We delivered a reproducible, validator-clean entity-resolution submission using conservative normalized exact matching, country partitioning, and global exclusivity.'); set_font(r)
h1(doc, 'Executive summary')
add_text(doc, 'The challenge required linking noisy business records from Sources 2 and 3 to a deduplicated Source 1 reference set. We built an auditable matching pipeline, ran large-scale data and submission work on AWS infrastructure, and packaged the code, notebooks, outputs, and documentation for reproduction.')
add_table(doc, ['Measured result', 'Value'], [
    ['Held-out macro F₀.₅', '0.141762'],
    ['Baseline macro F₀.₅', '0.125270'],
    ['Improvement', '+0.016492'],
    ['Required Source 1 test rows', '1,732,544'],
    ['Non-empty final predictions', '218,189'],
    ['Submission validation', 'PASS'],
], [3600, 5760])
add_text(doc, 'Important accuracy note: the shipped solution is a deterministic rule-based matcher; it does not claim that a neural model was trained. SageMaker-based GBM and neural reranking are documented as the scalable next-stage architecture and training plan.')

page_break(doc)
h1(doc, '1. Challenge, data, and solution delivered')
h2(doc, 'Problem statement')
add_text(doc, 'This is a one-to-many business entity-resolution task. A Source 1 business can correspond to zero or more noisy records in Sources 2 and 3. The records have no shared identifier and contain variations in business name, address, spelling, punctuation, legal suffixes, abbreviations, language/script, and token order.')
h2(doc, 'Data scale and measured findings')
add_table(doc, ['Dataset', 'Source 1', 'Source 2', 'Source 3'], [
    ['Training rows', '2,206,821', '5,034,616', '5,285,603'],
    ['Test rows', '1,732,544', '4,887,273', '5,082,316'],
], [1800, 2520, 2520, 2520])
add_bullet(doc, '7,638,365 labelled positive training pairs; average 3.46 matches per Source 1 entity.')
add_bullet(doc, 'No duplicate matched Source 2/3 IDs were found in the training ground truth, supporting a strict exclusivity constraint.')
add_bullet(doc, 'No cross-country matches were observed in the sampled ground truth, so matching was country-partitioned without hard-coding country labels.')
add_bullet(doc, '14.9% of true pairs shared no name token, but the sampled true pairs always shared a name or address token; this supports a name-plus-address retrieval strategy.')
h2(doc, 'Final submission method')
for x in [
    'Normalize business names: lowercasing, punctuation handling, Latin-diacritic folding, placeholder removal, organizer-provided legal-suffix handling, and &/and equivalence.',
    'Normalize addresses: tokenization, placeholder removal, order-invariant keys, and organizer-provided Rd/Road and St/Street canonicalization.',
    'Generate an exact joint key: (country, normalized name key, normalized address key).',
    'Accept only exact joint-key candidates and apply global exclusivity so an S2/S3 ID is assigned at most once.',
    'Write matching_results.tsv and candidate_pairs.tsv, then verify them with the organizer-supplied validator.'
]: add_number(doc, x)

page_break(doc)
h1(doc, '2. AWS work completed and reproducible workflow')
h2(doc, 'Amazon SageMaker AI')
add_text(doc, 'SageMaker AI was used as the notebook-oriented ML execution environment and remains the documented alternative for scale-up. The project is portable between local development, SageMaker Studio JupyterLab, and a SageMaker Notebook Instance through a common configuration cell that detects the runtime and assigns data, artifact, and output paths.')
add_table(doc, ['SageMaker activity', 'What to say / show'], [
    ['Notebook environment', 'JupyterLab notebooks for setup, metrics/splits, normalization, blocking, model experiments, decision layer, and submission generation.'],
    ['Reproducibility', 'Pinned requirements, portable paths, Parquet intermediates, notebook stages, and run logs in artifacts/runs.jsonl.'],
    ['Compute strategy', 'Small CPU for setup/EDA; high-memory CPU for blocking/features/LightGBM; GPU only for embedding or neural reranker experiments.'],
    ['Long jobs', 'Terminal-based nohup execution or SageMaker Notebook Scheduler for jobs that should survive browser closure.'],
], [2600, 6760])
h2(doc, 'EC2 terminal work')
add_text(doc, 'EC2 became the primary execution runtime for the challenge, using the ber-runner environment under ~/ber with Python 3.12. The terminal was used for environment setup, running notebooks/pipeline commands, validating outputs, viewing logs, and executing memory-bounded mining and blocking pilots.')
add_bullet(doc, 'Use the terminal to demonstrate: repository structure, Python version, dependency installation, output files, artifact logs, and validator command/result.')
add_bullet(doc, 'State the practical constraint: the 8 GiB EC2 machine is suitable for memory-bounded S3 mining and blocking pilots; full feature/model stages need larger CPU memory or GPU capacity.')
h2(doc, 'AWS Console and Amazon S3')
add_text(doc, 'The AWS Console is the control plane for SageMaker domains/spaces, EC2 instances, IAM roles, quotas, storage, and cost visibility. Amazon S3 is the durable data-staging layer: raw challenge data can be uploaded from the local workspace to an S3 bucket, then synchronized into SageMaker/EC2 compute close to the runtime.')
add_table(doc, ['S3 object group', 'Purpose'], [
    ['ber/dataset/', 'Training and test TSV data staged for compute.'],
    ['artifacts/', 'Parquet intermediate outputs, splits, alias tables, and experiment logs.'],
    ['output/', 'Final matching_results.tsv and candidate_pairs.tsv.'],
    ['submission package/', 'Code, notebooks, README, requirements, and documentation for audit/reproduction.'],
], [3000, 6360])

page_break(doc)
h1(doc, '3. Model-training vocabulary and scalable ML architecture')
add_text(doc, 'Use this section in a presentation to explain the full ML design without overstating the final implementation. The rule-based final model is the baseline; the following cascade is the planned/experiment-ready training architecture.')
h2(doc, 'Proposed cascade')
add_table(doc, ['Stage', 'Purpose', 'Key terms'], [
    ['1. Blocking / retrieval', 'Reduce billions of possible record pairs to a manageable candidate set while preserving true matches.', 'candidate generation, inverted index, TF-IDF, char n-grams, recall@K, reduction ratio'],
    ['2. GBM scorer', 'Score candidate pairs using structured similarity features.', 'LightGBM/gradient boosting, features, class imbalance, calibration, feature ablation'],
    ['3. Neural reranker', 'Apply an expensive semantic model only to the highest-value candidates.', 'embeddings, bi-encoder, cross-encoder, transformer, top-K reranking, GPU inference'],
    ['4. Decision layer', 'Convert pair scores into final record links under business rules.', 'threshold tuning, macro F₀.₅ optimization, exclusivity-constrained assignment, precision/recall trade-off'],
], [1700, 3300, 4360])
h2(doc, 'Training and evaluation terms')
add_table(doc, ['Term', 'Meaning in this project'], [
    ['Train / validation split', 'A country × match-cardinality stratified holdout isolates evaluation entities from training entities.'],
    ['Ground truth / labels', 'Known Source 1 → Source 2/3 entity links used to evaluate or train the system.'],
    ['Positive / negative pair', 'A correct matching pair versus a non-matching candidate pair; negatives are much more common.'],
    ['Features', 'Name/address token overlap, string similarity, numeric agreement, country agreement, and raw vs. normalized comparisons.'],
    ['F₀.₅ score', 'Macro-averaged metric that weights precision more than recall; threshold must be tuned empirically.'],
    ['Recall@K', 'Share of true matches retained among the top K candidates; blocking quality ceiling.'],
    ['Ablation study', 'Remove one feature group or normalization step at a time to measure its contribution.'],
    ['Hyperparameter sweep', 'Systematic search over K, thresholds, GBM settings, normalization variants, and reranker choices.'],
    ['Overfitting', 'A model improves on training data but not on the held-out split; controlled through a fixed validation protocol.'],
], [2300, 7060])

page_break(doc)
h1(doc, '4. System requirements, tooling, and demonstration checklist')
h2(doc, 'Technical requirements')
add_table(doc, ['Area', 'Requirement / recommended configuration'], [
    ['Operating environment', 'Python 3.12; Linux-based EC2 or SageMaker Studio JupyterLab runtime.'],
    ['Core packages', 'numpy 2.1.3, pandas 2.2.3, pyarrow 17.0.0, duckdb 1.1.3, rapidfuzz 3.10.1, scikit-learn 1.5.2, scipy 1.14.1, lightgbm 4.5.0, jupyterlab 4.3.1, ipykernel 6.29.5.'],
    ['Storage', 'Approximately 300 GB gp3 for raw data, Parquet intermediates, outputs, checkpoints, and operating buffer.'],
    ['CPU workloads', '16 vCPU and 64 GB+ RAM recommended for blocking, features, and LightGBM at full scale (ml.r7i.4xlarge / ml.m7i.4xlarge class).'],
    ['GPU workloads', 'For neural embeddings/reranking: an A10G-class GPU with 24 GB VRAM (ml.g5.2xlarge class), subject to quota.'],
    ['Data formats', 'TSV for source/submission files; Parquet for intermediate artifacts; JSONL for experiment logs.'],
    ['Security / access', 'Least-privilege IAM role for SageMaker/EC2 to read/write the approved S3 bucket; no external entity databases or lookup APIs.'],
], [2350, 7010])
h2(doc, 'What to show to judges')
for x in [
    'Problem and scale: explain why naive all-pairs matching is infeasible and why candidate blocking is necessary.',
    'Architecture: show the lexical blocking → scorer → reranker → exclusivity decision flow, clearly noting which layers are final versus planned.',
    'AWS Console: SageMaker/EC2 compute selection, S3 bucket structure, IAM-aware data access, and cost-conscious stop/start practice.',
    'Terminal: portable environment, requirements file, project notebooks/source tree, and a validation command or saved PASS result.',
    'Data and rules: no external enrichment; organizer-provided normalizations only in the shipped method.',
    'Results: F₀.₅ 0.141762 versus 0.125270 baseline; 1,732,544 required rows and validator PASS.',
    'Reproducibility: code package, notebooks, configuration, Parquet artifacts, output TSVs, and run log.'
]: add_bullet(doc, x)
h2(doc, 'High-value keywords')
add_text(doc, 'Amazon SageMaker AI, SageMaker Studio, JupyterLab, Amazon EC2, Amazon S3, AWS Console, IAM role, cloud storage, ML pipeline, entity resolution, record linkage, data normalization, blocking, candidate generation, lexical retrieval, TF-IDF, n-grams, string similarity, RapidFuzz, feature engineering, LightGBM, gradient boosting, embeddings, transformer, reranking, threshold tuning, stratified holdout, macro F₀.₅, precision, recall, recall@K, exclusivity constraint, reproducibility, experiment tracking, Parquet, JSONL, validation, scalable inference, GPU acceleration, cost optimization.')

page_break(doc)
h1(doc, '5. Evidence, deliverables, and next steps')
h2(doc, 'Delivered artefacts')
add_table(doc, ['Artefact', 'Evidence / use'], [
    ['Source code', 'Normalization, blocking, metric, split, submission, and rule-baseline modules under src/business_entity_resolution/.'],
    ['Notebooks', 'NB00 setup; NB02 metric and split; NB03 normalization; NB04 blocking; NB09 final submission, with later-stage notebook plan documented.'],
    ['Outputs', 'matching_results.tsv and candidate_pairs.tsv produced for the required test set.'],
    ['Validation', 'Organizer validator passed with all 1,732,544 Source 1 rows present and no format issues.'],
    ['Documentation', 'Project plan, decision log, SageMaker runbook, README, requirements, and final submission package.'],
], [2400, 6960])
h2(doc, 'Suggested next improvements')
add_bullet(doc, 'Measure lexical union blocking recall@K on the full validation split and tune K within compute budget.')
add_bullet(doc, 'Generate feature tables for candidates and train/compare LightGBM configurations using the fixed stratified holdout.')
add_bullet(doc, 'Add a neural reranker only if it provides a measured gain over the GBM stage and is viable within GPU quota/cost limits.')
add_bullet(doc, 'Tune global and segmented acceptance thresholds directly against macro F₀.₅, then analyze errors by country, source, cardinality, and corruption type.')
add_bullet(doc, 'Record each experiment configuration, runtime, validation metrics, and validator outcome in artifacts/runs.jsonl.')
h2(doc, 'Short presentation script')
add_text(doc, '“Our challenge was to connect noisy business records at millions-of-rows scale. We first made the process reliable: portable notebooks, S3-based staging, reproducible artifacts, and a validator-clean submission. Our final system is a precise, explainable joint-key matcher with country partitioning and global exclusivity, improving held-out macro F₀.₅ from 0.125270 to 0.141762. The scalable next layer is a candidate-retrieval, LightGBM, and neural-reranking cascade on SageMaker, where expensive models only score a small top-K candidate set.”')
h2(doc, 'Source basis')
add_text(doc, 'This report is based on the project plan, decision log, SageMaker runbook, pinned requirements, executed submission documentation, source modules, and final output artefacts in the challenge workspace. Values marked as final refer to the documented held-out evaluation and validator result.')

doc.save(OUT)
print(OUT)
